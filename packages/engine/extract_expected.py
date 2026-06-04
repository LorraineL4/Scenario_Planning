#!/usr/bin/env python3
"""
extract_expected.py  —  Extract computed output values from DAP Excel PG sheets.

Reads the calculated cells from each Product Group sheet and writes them to
expected.json. Used as ground-truth for the engine test suite.

Usage:
    python packages/engine/extract_expected.py "data/raw/260413 - WFM_CORP - DAP.xlsx"
    python packages/engine/extract_expected.py "data/raw/260413 - WFM_CORP - DAP.xlsx" --output data/processed/expected.json

Output JSON shape:
  {
    "skus": {
      "<sku_name>": {
        "P01": { unit_sales, baseline_weekly, promo_units, ... },
        ...
        "P12": { ... }
      }
    }
  }

NOTE: Output contains client data — do not commit to git.
      Recommended: write to data/processed/ (gitignored).
"""

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl not installed.  Run:  pip install openpyxl")

from extract_config import PERIODS, PERIOD_COL, _num, _is_pg_sheet, load_fiscal_calendar, load_pg_list, _col_letter_to_num

# Maps engine output field name → PG sheet row number
OUTPUT_ROWS = {
    "baseline_weekly":            100,
    "unit_sales":                  97,
    "promo_units":                101,
    "non_promo_units":            102,
    "incremental_units":          103,
    "base_units":                 104,
    "gross_sales":                108,
    "total_cogs":                 112,
    "retailer_working_spend":     116,
    "distributor_working":        194,
    "digital_sales_spend":        195,
    "other_spend":                196,
    "terms_spoils_spend":         197,
    "total_spend":                200,
    "profit_after_working_spend": 128,
    "profit_after_total_spend":   201,
}


def extract_pg_expected(ws, period_col=None) -> dict:
    """Extract per-period expected outputs from a PG sheet."""
    if period_col is None:
        period_col = PERIOD_COL
    periods = {}
    for period, col in period_col.items():
        period_data = {}
        for field, row in OUTPUT_ROWS.items():
            period_data[field] = _num(ws.cell(row=row, column=col).value)
        periods[period] = period_data
    return periods


def extract(excel_path: Path) -> dict:
    print(f"Loading {excel_path.name} ...", flush=True)
    wb = openpyxl.load_workbook(str(excel_path), data_only=True)

    fiscal_calendar = load_fiscal_calendar(wb)
    period_col = {p: _col_letter_to_num(d["col"]) for p, d in fiscal_calendar.items()}
    pg_list    = load_pg_list(wb)

    pg_names = [s for s in wb.sheetnames if _is_pg_sheet(wb, s, pg_names=pg_list)]
    print(f"  [+] Product Group sheets ({len(pg_names)} found):", flush=True)

    skus = {}
    for name in pg_names:
        print(f"      - {name}", flush=True)
        skus[name] = extract_pg_expected(wb[name], period_col=period_col)

    return {
        "_meta": {
            "source_file":  excel_path.name,
            "extracted_at": datetime.now(timezone.utc).isoformat(),
            "output_rows":  OUTPUT_ROWS,
        },
        "skus": skus,
    }


def main():
    parser = argparse.ArgumentParser(
        description="Extract DAP computed output values to JSON (ground truth for tests).",
        epilog="Output file contains client data — do not commit to git.",
    )
    parser.add_argument("excel_path", help="Path to DAP .xlsx file")
    parser.add_argument(
        "--output", "-o",
        default="data/processed/expected.json",
        help="Output JSON path (default: data/processed/expected.json)",
    )
    args = parser.parse_args()

    path = Path(args.excel_path)
    if not path.exists():
        sys.exit(f"File not found: {path}")

    data = extract(path)

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, indent=2, default=str), encoding="utf-8")

    n_skus = len(data["skus"])
    print(f"\nWrote {out}  ({out.stat().st_size:,} bytes)")
    print(f"  {n_skus} SKU(s)  |  {len(PERIODS)} periods  |  {len(OUTPUT_ROWS)} fields each")


if __name__ == "__main__":
    main()
